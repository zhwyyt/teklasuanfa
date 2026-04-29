from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


FAMILY_ROW_COLUMNS = [
    ("MemberId", "构件编号"),
    ("AssemblyId", "装配编号"),
    ("DefinitionDrivenFamilyLabelZh", "定义驱动家族"),
    ("LongitudinalTypeLabelZh", "长度方向类型"),
    ("LongitudinalSubtypeLabelZh", "长度方向细分类"),
    ("DefinitionDrivenSubtypeLabelZh", "定义驱动子类"),
    ("DecisionStatusLabelZh", "判定状态"),
    ("DecisionReasonLabelZh", "判定原因"),
    ("LeadClauseLabelZh", "主条款"),
    ("LeadClauseVerdictLabelZh", "主条款判定"),
    ("LeadClausePromotionReadinessLabelZh", "提升准备度"),
    ("LeadClauseEffectDirectionLabelZh", "effect 方向"),
    ("LeadClauseBreakEffectCount", "破坏 effect 数"),
    ("LeadClauseRewriteEffectCount", "改写 effect 数"),
    ("BodyDescriptorFamily", "主体描述家族"),
    ("BodyDescriptorSectionType", "主体描述截面类型"),
    ("SourceSemanticBodyFamily", "源语义主体家族"),
    ("SourceSemanticSectionType", "源语义截面类型"),
    ("SatisfiedConditions", "已满足条件"),
    ("MissingConditions", "缺失条件"),
    ("SourceFile", "源文件"),
]

PROFILE_ROW_COLUMNS = [
    ("MemberId", "构件编号"),
    ("AssemblyId", "装配编号"),
    ("FamilyLabelZh", "家族"),
    ("LongitudinalTypeLabelZh", "长度方向类型"),
    ("LongitudinalSubtypeLabelZh", "长度方向细分类"),
    ("FamilySubtypeLabelZh", "家族子类"),
    ("ResolutionStatusLabelZh", "细分状态"),
    ("ProfileCategoryLabelZh", "型材大类"),
    ("ProfileLabelZh", "型材"),
    ("ProfileSeriesLabelZh", "型材系列"),
    ("DimensionSourceLabelZh", "尺寸来源"),
    ("SimilarityBasisLabelZh", "命中依据"),
    ("SimilarityScore", "相似度"),
    ("EvidenceSummary", "证据摘要"),
    ("SatisfiedConditions", "已满足条件"),
    ("MissingConditions", "缺失条件"),
    ("SourceFile", "源文件"),
]

COARSE_MAIN_CLASS_ROW_COLUMNS = [
    ("MemberId", "构件编号"),
    ("AssemblyId", "装配编号"),
    ("SourceMemberMainClassLabelZh", "属性直达主类"),
    ("BodyDescriptorFamily", "主体描述家族"),
    ("BodyDescriptorSectionType", "主体描述截面类型"),
    ("LongitudinalTypeLabelZh", "长度方向类型"),
    ("CoarseMainClassLabelZh", "粗主类"),
    ("CoarseMainClassSubtypeLabelZh", "粗子类"),
    ("CoarseMainClassConfidence", "粗分类置信度"),
    ("CandidatePartCount", "主板候选数量"),
    ("CandidatePartIds", "主板候选零件"),
    ("EligibleStationCount", "有效切片数"),
    ("BoxStationCount", "箱体切片数"),
    ("HStationCount", "H类切片数"),
    ("PrimaryPlateStationCount", "主板切片数"),
    ("ClosedLoopStationCount", "闭环切片数"),
    ("BoxStationRatio", "箱体切片占比"),
    ("HStationRatio", "H类切片占比"),
    ("PrimaryPlateStationRatio", "主板切片占比"),
    ("ClosedLoopStationRatio", "闭环切片占比"),
    ("CoarseMainClassReasonLabelZh", "粗分类原因"),
    ("SourceFile", "源文件"),
]

SUMMARY_SECTIONS = [
    ("StatusBreakdown", "状态分布"),
    ("FamilyBreakdown", "家族分布"),
    ("LongitudinalTypeBreakdown", "长度方向类型分布"),
    ("SubtypeBreakdown", "子类分布"),
    ("DecisionReasonBreakdown", "判定原因分布"),
    ("ReasonBreakdown", "判定原因分布"),
    ("CategoryBreakdown", "型材大类分布"),
    ("ProfileBreakdown", "型材分布"),
    ("ProfileSeriesBreakdown", "型材系列分布"),
    ("DimensionSourceBreakdown", "尺寸来源分布"),
]


def _load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def _stringify(value: object) -> object:
    if isinstance(value, list):
        return "\n".join(str(item) for item in value)
    return value


def _autosize(worksheet) -> None:
    for column_cells in worksheet.columns:
        max_length = 0
        column_index = column_cells[0].column
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 10), 48)


def _style_header(row) -> None:
    fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    font = Font(bold=True)
    for cell in row:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _add_summary_sheet(workbook: Workbook, title: str, summary: dict) -> None:
    worksheet = workbook.active
    worksheet.title = title
    worksheet.append(["字段", "值"])
    _style_header(worksheet[1])

    for key in ("AssemblyCount", "AdjudicatedCount", "DeferredCount", "ResolvedCount", "ReviewBypassCount"):
        if key in summary:
            worksheet.append([key, summary[key]])

    for section_key, section_label in SUMMARY_SECTIONS:
        rows = summary.get(section_key)
        if not rows:
            continue
        worksheet.append([])
        worksheet.append([section_label, ""])
        _style_header(worksheet[worksheet.max_row])
        worksheet.append(["Code", "LabelZh", "Count"])
        _style_header(worksheet[worksheet.max_row])
        for row in rows:
            worksheet.append([row.get("Code", ""), row.get("LabelZh", ""), row.get("Count", "")])

    worksheet.freeze_panes = "A2"
    _autosize(worksheet)


def _add_rows_sheet(workbook: Workbook, title: str, rows: Iterable[dict], columns: list[tuple[str, str]]) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.append([label for _, label in columns])
    _style_header(worksheet[1])

    for row in rows:
        worksheet.append([_stringify(row.get(key, "")) for key, _ in columns])

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    _autosize(worksheet)


def _export_workbook(source_path: Path, target_path: Path, summary_sheet_title: str, row_sheet_title: str, row_columns):
    payload = _load_json(source_path)
    workbook = Workbook()
    _add_summary_sheet(workbook, summary_sheet_title, payload.get("Summary", {}))
    _add_rows_sheet(workbook, row_sheet_title, payload.get("Rows", []), row_columns)
    workbook.save(target_path)


def main() -> int:
    if len(sys.argv) != 2:
        print("Usage: python tools/Export-RecognitionExcels.py <output-directory>")
        return 1

    output_dir = Path(sys.argv[1]).resolve()
    if not output_dir.is_dir():
        print(f"Output directory not found: {output_dir}")
        return 1

    family_json = output_dir / "body-family-proof.json"
    profile_json = output_dir / "body-profile-resolution.json"
    coarse_json = output_dir / "coarse-main-class-observation.json"

    if not family_json.exists() or not profile_json.exists():
        print("Required JSON artifacts are missing. Expecting body-family-proof.json and body-profile-resolution.json.")
        return 1

    family_xlsx = output_dir / "body-family-proof.xlsx"
    profile_xlsx = output_dir / "body-profile-resolution.xlsx"
    coarse_xlsx = output_dir / "coarse-main-class-observation.xlsx"

    _export_workbook(family_json, family_xlsx, "家族判定汇总", "家族判定明细", FAMILY_ROW_COLUMNS)
    _export_workbook(profile_json, profile_xlsx, "型材细分汇总", "型材细分明细", PROFILE_ROW_COLUMNS)
    if coarse_json.exists():
        _export_workbook(coarse_json, coarse_xlsx, "粗分类汇总", "粗分类明细", COARSE_MAIN_CLASS_ROW_COLUMNS)

    print(f"Generated: {family_xlsx}")
    print(f"Generated: {profile_xlsx}")
    if coarse_json.exists():
        print(f"Generated: {coarse_xlsx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
